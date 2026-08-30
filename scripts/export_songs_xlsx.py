#!/usr/bin/env python3
"""Export the song list from index.html into a simple spreadsheet.

Reads the window.TRACKS block in index.html (the "EDIT YOUR TRACKS HERE"
section) and writes one row per track with the song name, its Spotify link,
the button/category it belongs to, and the start time.

Usage:
    python scripts/export_songs_xlsx.py [output.xlsx]
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tracks_io import (
    CATEGORY_LABELS,
    INDEX_HTML,
    REPO_ROOT,
    extract_tracks,
    format_start_time,
    spotify_url,
)

DEFAULT_OUTPUT = REPO_ROOT / "wild-hockey-songs.xlsx"

HEADERS = ["Song Name", "Spotify Link", "Category", "Start Time"]


def build_rows(tracks: dict) -> list:
    rows = []
    for key, entries in tracks.items():
        label = CATEGORY_LABELS.get(key, key)
        for entry in entries:
            rows.append(
                [
                    entry.get("name", ""),
                    spotify_url(entry.get("uri", "")),
                    label,
                    format_start_time(entry.get("startSec")),
                ]
            )
    return rows


def write_workbook(rows: list, output: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Songs"

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="154734")  # Wild green
    body_font = Font(name="Arial")
    link_font = Font(name="Arial", color="0563C1", underline="single")

    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")

    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = body_font
        link_cell = ws.cell(row=r, column=2)
        if link_cell.value:
            link_cell.hyperlink = link_cell.value
            link_cell.font = link_font
        ws.cell(row=r, column=4).alignment = Alignment(horizontal="left")

    note_row = len(rows) + 3
    note = ws.cell(
        row=note_row,
        column=1,
        value=(
            "Exported from the window.TRACKS block in index.html. Start Time is shown as "
            "m:ss and corresponds to the startSec value (in seconds); a blank means the "
            "track plays from the beginning."
        ),
    )
    note.font = Font(name="Arial", italic=True, size=9)

    widths = [30, 52, 22, 12]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{len(rows) + 1}"

    wb.save(output)


def main() -> None:
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUTPUT
    tracks = extract_tracks(INDEX_HTML.read_text(encoding="utf-8"))
    rows = build_rows(tracks)
    write_workbook(rows, output)
    print(f"Wrote {len(rows)} songs to {output}")


if __name__ == "__main__":
    main()
